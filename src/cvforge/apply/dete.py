"""Irish employment permit statistics (DETE): which companies were issued permits."""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from urllib.parse import urljoin

from cvforge.apply import http
from cvforge.apply.companies import _save, load_registry, name_key
from cvforge.apply.http import FetchError, host_of, on_domain

STATS_INDEX = "https://enterprise.gov.ie/en/what-we-do/workplace-and-skills/employment-permits/statistics/"
DETE_DOMAIN = "enterprise.gov.ie"
IE_SUFFIXES = frozenset({"limited", "unlimited company", "dac", "designated activity company",
                         "ireland", "europe", "emea", "uc", "clg"})
MARK_TYPE = "ie-permit"
HREF = re.compile(r"""href\s*=\s*["']([^"']+)["']""", re.IGNORECASE)
YEAR_PAGE = re.compile(r"employment-permit-statistics-(\d{4})", re.IGNORECASE)
PREFIX_MIN = 4
GENERIC = frozenset({
    "irish", "dublin", "cork", "galway", "limerick", "kerry", "global", "international",
    "national", "tech", "data", "cloud", "health", "care", "smart", "digital", "first", "new",
    "the", "group", "services", "solutions", "systems"})
SOURCE_YEAR = re.compile(r"\b(20\d{2})\b")
# Words that may follow a registry name in a DETE name for a prefix match to count:
# entity descriptors only ("Stripe Payments", "Okta Identity"), never another business
# ("Railway & Steam Packet Union", "Citi Bus"). Phrases are in dete_key form.
DESCRIPTORS = frozenset({
    "payments", "technology", "technologies", "software", "labs", "holdings", "international",
    "identity", "r d", "ai", "operations", "financial", "services", "sistemas", "s a",
    "irish branch", "hibernia"})
TOTAL_ROWS = frozenset({"grand total", "total"})


class DeteError(Exception):
    """The DETE data could not be read."""


class NoYearPage(FetchError):
    """DETE has not published a statistics page for the year yet."""


@dataclass
class ImportResult:
    year: int
    employers: int
    matched: list[str] = field(default_factory=list)
    details: dict[str, list[str]] = field(default_factory=dict)  # company -> DETE names
    possible: dict[str, list[str]] = field(default_factory=dict)  # unconfirmed prefix names


def dete_key(name: str) -> str:
    return name_key(name, IE_SUFFIXES)


def year_page_links(index_html: str) -> dict[int, str]:
    out: dict[int, str] = {}
    for href in HREF.findall(index_html):
        m = YEAR_PAGE.search(href)
        if m and href.lower().split("?")[0].endswith((".html", "/")):
            out.setdefault(int(m.group(1)), href)
    return out


def xlsx_link(year_html: str) -> str | None:
    for href in HREF.findall(year_html):
        low = href.lower().split("?")[0]
        if "permits-issued-to-companies" in low and low.endswith(".xlsx"):
            return href
    return None


def _is_total_header(cell: str) -> bool:
    return cell.startswith(("grand total", "permits issued grand total"))


def _header(row) -> tuple[int, int] | None:
    """(name column, total column) when the row is a header row, else None.

    2026 layout: "Employer Name" ... "Permits Issued Grand Total". 2025 layout: an empty
    name header, months, then "Grand Total".
    """
    cells = [str(c).strip().lower() if c is not None else "" for c in row]
    total = next((i for i, c in enumerate(cells) if _is_total_header(c)), None)
    if total is None:
        return None
    name = cells.index("employer name") if "employer name" in cells else 0
    return (name, total) if name != total else None


def _parse_sheet(ws) -> dict[str, int] | None:
    """Permits per employer, or None when the sheet has no header row."""
    cols = None
    out: dict[str, int] = {}
    for row in ws.iter_rows(values_only=True):
        if cols is None:
            cols = _header(row)
            continue
        name_col, total_col = cols
        if len(row) <= max(name_col, total_col) or row[name_col] is None:
            continue
        name = str(row[name_col]).strip()
        if not name or name.lower() in TOTAL_ROWS:
            continue
        try:
            out[name] = out.get(name, 0) + int(row[total_col] or 0)
        except (TypeError, ValueError):
            continue
    return out if cols is not None else None


def parse_company_xlsx(data: bytes) -> dict[str, int]:
    """The "Export" sheet when present, else the first sheet with a header row."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types on bad input
        raise DeteError(f"unreadable spreadsheet: {exc}") from exc
    sheets = ["Export"] if "Export" in wb.sheetnames else wb.sheetnames
    for sheet in sheets:
        found = _parse_sheet(wb[sheet])
        if found is not None:
            return found
    raise DeteError("no header row with an employer name and a 'Grand Total' column")


def _descriptors_only(rest: str) -> bool:
    """True when every word of `rest` is an entity descriptor (two-word phrases first)."""
    words = rest.split()
    i = 0
    while i < len(words):
        if " ".join(words[i:i + 2]) in DESCRIPTORS and i + 1 < len(words):
            i += 2
        elif words[i] in DESCRIPTORS:
            i += 1
        else:
            return False
    return bool(words)


def _prefixed(key: str, employers: dict[str, int]) -> tuple[list[str], list[str], list[str]]:
    """(exact, confirmed prefix, possible prefix) DETE names for a company key."""
    exact: list[str] = []
    prefix: list[str] = []
    possible: list[str] = []
    if not key:
        return exact, prefix, possible
    allow_prefix = len(key) >= PREFIX_MIN and key not in GENERIC
    for name in employers:
        ekey = dete_key(name)
        if ekey == key:
            exact.append(name)
        elif allow_prefix and ekey.startswith(key + " "):
            (prefix if _descriptors_only(ekey[len(key) + 1:]) else possible).append(name)
    return exact, prefix, possible


def _matches(key: str, employers: dict[str, int]) -> tuple[int, list[str]] | None:
    """(permit total, DETE names) for a company key.

    An exact DETE name wins. Otherwise DETE names starting with the key plus a space, but only
    for a distinctive key (4+ chars, not a generic word) and only when the rest of the DETE
    name is entity descriptors ("Stripe Payments Europe Limited", not "Citi Bus Limited").
    """
    exact, prefix, _ = _prefixed(key, employers)
    names = exact or prefix
    return (sum(employers[n] for n in names), names) if names else None


def _cache(data_dir: Path, year: int) -> Path:
    return Path(data_dir) / "sponsors" / f"ie-{year}.json"


def _source(year: int) -> str:
    return f"DETE employment permits issued to companies {year}"


def _fetch_dete(fetch, url: str) -> bytes:
    if not on_domain(host_of(url), DETE_DOMAIN):
        raise FetchError(f"not a {DETE_DOMAIN} URL: {url}")
    return fetch(url)


def import_year(data_dir: Path, year: int, *, fetch=http.get_bytes, file: Path | None = None,
                index_url: str = STATS_INDEX) -> ImportResult:
    data_dir = Path(data_dir)
    if file is not None:
        data, source = Path(file).read_bytes(), str(Path(file).name)
    else:
        index = _fetch_dete(fetch, index_url).decode("utf-8", errors="replace")
        page = year_page_links(index).get(year)
        if page is None:
            raise NoYearPage(f"no DETE statistics page for {year}")
        page_url = urljoin(index_url, page)
        link = xlsx_link(_fetch_dete(fetch, page_url).decode("utf-8", errors="replace"))
        if link is None:
            raise FetchError(f"no companies spreadsheet on {page_url}")
        source = urljoin(page_url, link)
        data = _fetch_dete(fetch, source)
    employers = parse_company_xlsx(data)
    cache = _cache(data_dir, year)
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps({"year": year, "source": source,
                                 "fetched": date.today().isoformat(), "employers": employers},
                                ensure_ascii=False), encoding="utf-8")
    details, possible = _mark_registry(data_dir / "companies.yaml", year, employers)
    return ImportResult(year, len(employers), list(details), details, possible)


def _mark_registry(path: Path, year: int, employers: dict[str, int]
                   ) -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    """(matched DETE names, possible DETE names) per registry company; marks the matches."""
    registry = load_registry(path)
    details: dict[str, list[str]] = {}
    possible: dict[str, list[str]] = {}
    today = date.today().isoformat()
    source = _source(year)
    changed = False
    for entry in registry:
        hits = [h for n in [entry["name"], *entry.get("aliases", [])]
                if (h := _matches(dete_key(n), employers)) is not None]
        kept = [s for s in entry["sponsor"] if not (s["type"] == MARK_TYPE and s["source"] == source)]
        if hits:
            count, names = max(hits, key=lambda h: h[0])
            kept.append({"type": MARK_TYPE, "source": source, "date": today, "count": count,
                         "matched": sorted(names)})
            details[entry["name"]] = sorted(names)
        else:
            maybe = {p for n in [entry["name"], *entry.get("aliases", [])]
                     for p in _prefixed(dete_key(n), employers)[2]}
            if maybe:
                possible[entry["name"]] = sorted(maybe)
        if kept != entry["sponsor"] or hits:
            changed = True
        entry["sponsor"] = kept
    if changed:
        _save(path, registry)
    return details, possible


def _cached_years(data_dir: Path) -> list[tuple[int, Path]]:
    found = []
    for p in (Path(data_dir) / "sponsors").glob("ie-*.json"):
        m = re.fullmatch(r"ie-(\d{4})\.json", p.name)
        if m:
            found.append((int(m.group(1)), p))
    return sorted(found, reverse=True)


def _read_cache(path: Path) -> dict[str, int] | None:
    try:
        employers = json.loads(path.read_text(encoding="utf-8"))["employers"]
        return employers if isinstance(employers, dict) else None
    except (OSError, ValueError, KeyError, TypeError):
        return None


def dete_imported(data_dir: Path) -> bool:
    return any(_read_cache(p) is not None for _, p in _cached_years(data_dir))


def permit_record(data_dir: Path, company_name: str, registry_entry: dict | None
                  ) -> tuple[int, int] | None:
    """(year, count) for the newest year with a match, from the registry mark or the cache."""
    best: tuple[int, int] | None = None
    for s in (registry_entry or {}).get("sponsor", []):
        if s.get("type") == MARK_TYPE and "count" in s:
            m = SOURCE_YEAR.search(s.get("source", ""))
            if m and (best is None or int(m.group(1)) > best[0]):
                best = (int(m.group(1)), s["count"])
    key = dete_key(company_name)
    for year, path in _cached_years(data_dir):
        if best and year <= best[0]:
            break
        employers = _read_cache(path)
        hit = _matches(key, employers) if employers else None
        if hit is not None:
            return year, hit[0]
    return best
